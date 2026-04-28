from __future__ import annotations

import threading
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

from centralizacion_remuneraciones import run_pipeline


class FilePicker(ttk.LabelFrame):
    def __init__(self, master, title: str, multiple: bool = True):
        super().__init__(master, text=title, padding=8)
        self.multiple = multiple
        self.paths: list[str] = []

        self.listbox = tk.Listbox(self, height=4, width=110)
        self.listbox.grid(row=0, column=0, columnspan=3, sticky="nsew")
        ttk.Button(self, text="Seleccionar", command=self.select_files).grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Button(self, text="Quitar", command=self.remove_selected).grid(row=1, column=1, sticky="w", padx=8, pady=(8, 0))
        ttk.Button(self, text="Limpiar", command=self.clear).grid(row=1, column=2, sticky="w", pady=(8, 0))
        self.columnconfigure(0, weight=1)

    def select_files(self):
        if self.multiple:
            chosen = filedialog.askopenfilenames(filetypes=[("Excel", "*.xls *.xlsx")])
            if chosen:
                for path in chosen:
                    if path not in self.paths:
                        self.paths.append(path)
        else:
            chosen = filedialog.askopenfilename(filetypes=[("Archivos", "*.*")])
            if chosen:
                self.paths = [chosen]
        self.refresh()

    def refresh(self):
        self.listbox.delete(0, tk.END)
        for path in self.paths:
            self.listbox.insert(tk.END, path)

    def remove_selected(self):
        selected = set(self.listbox.curselection())
        self.paths = [p for i, p in enumerate(self.paths) if i not in selected]
        self.refresh()

    def clear(self):
        self.paths = []
        self.refresh()


class SinglePath(ttk.Frame):
    def __init__(self, master, title: str, filetypes, directory: bool = False):
        super().__init__(master)
        self.filetypes = filetypes
        self.directory = directory
        ttk.Label(self, text=title, width=28).grid(row=0, column=0, sticky="w")
        self.var = tk.StringVar()
        ttk.Entry(self, textvariable=self.var, width=95).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(self, text="Buscar", command=self.pick).grid(row=0, column=2, sticky="w")
        self.columnconfigure(1, weight=1)

    def pick(self):
        if self.directory:
            chosen = filedialog.askdirectory()
        else:
            chosen = filedialog.askopenfilename(filetypes=self.filetypes)
        if chosen:
            self.var.set(chosen)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Centralizacion de Remuneraciones")
        self.geometry("1100x860")
        self.month_options = [
            "ENERO",
            "FEBRERO",
            "MARZO",
            "ABRIL",
            "MAYO",
            "JUNIO",
            "JULIO",
            "AGOSTO",
            "SEPTIEMBRE",
            "OCTUBRE",
            "NOVIEMBRE",
            "DICIEMBRE",
        ]
        self.year_options = [str(year) for year in range(2025, 2036)]

        container = ttk.Frame(self, padding=12)
        container.pack(fill="both", expand=True)

        header = ttk.Label(
            container,
            text="Carga los archivos del mes, genera la centralizacion y completa cuentas nuevas si aparecen.",
        )
        header.pack(anchor="w", pady=(0, 10))

        top = ttk.Frame(container)
        top.pack(fill="x", pady=(0, 10))

        ttk.Label(top, text="Area").grid(row=0, column=0, sticky="w")
        self.area_var = tk.StringVar(value="Educacion")
        ttk.Combobox(top, textvariable=self.area_var, values=["Educacion", "Jardines"], width=18, state="readonly").grid(
            row=0, column=1, sticky="w", padx=(6, 16)
        )
        ttk.Label(top, text="Mes").grid(row=0, column=2, sticky="w")
        self.month_var = tk.StringVar(value="FEBRERO")
        ttk.Combobox(top, textvariable=self.month_var, values=self.month_options, width=16, state="readonly").grid(
            row=0, column=3, sticky="w", padx=(6, 16)
        )
        ttk.Label(top, text="Año").grid(row=0, column=4, sticky="w")
        self.year_var = tk.StringVar(value="2026")
        ttk.Combobox(top, textvariable=self.year_var, values=self.year_options, width=10, state="readonly").grid(
            row=0, column=5, sticky="w", padx=6
        )

        helper = ttk.Label(
            container,
            text="Identificación de proceso: si el nombre o carpeta contiene 'Proceso Normal' se toma como normal; si contiene 'Proceso 1', 'Proceso 2', etc., se agrupa con ese número.",
        )
        helper.pack(anchor="w", pady=(0, 8))

        self.maestros = FilePicker(container, "1. Maestro de Remuneraciones", multiple=True)
        self.maestros.pack(fill="x", pady=6)
        self.gastos = FilePicker(container, "2. InformeGastoFinanciamiento", multiple=True)
        self.gastos.pack(fill="x", pady=6)
        self.asientos = FilePicker(container, "3. Asiento Remuneracion (opcional, referencia)", multiple=True)
        self.asientos.pack(fill="x", pady=6)
        self.centrales = FilePicker(container, "4. Centralizacion CasChile actual (opcional)", multiple=True)
        self.centrales.pack(fill="x", pady=6)

        paths_frame = ttk.LabelFrame(container, text="Archivos auxiliares", padding=8)
        paths_frame.pack(fill="x", pady=8)
        self.mapeo_excel = SinglePath(paths_frame, "Mapeo Excel", [("Excel", "*.xlsx")])
        self.mapeo_excel.grid(row=0, column=0, sticky="ew", pady=4)
        self.salida_dir = SinglePath(paths_frame, "Carpeta salida", [], directory=True)
        self.salida_dir.grid(row=1, column=0, sticky="ew", pady=4)
        paths_frame.columnconfigure(0, weight=1)

        actions = ttk.Frame(container)
        actions.pack(fill="x", pady=(12, 8))
        self.generate_btn = ttk.Button(actions, text="Generar centralizacion", command=self.generate)
        self.generate_btn.pack(side="left")
        ttk.Button(actions, text="Cerrar", command=self.destroy).pack(side="right")

        self.status_var = tk.StringVar(value="Listo para cargar archivos.")
        ttk.Label(container, textvariable=self.status_var).pack(anchor="w", pady=(0, 6))

        self.log = tk.Text(container, height=14)
        self.log.pack(fill="both", expand=True)

        default_output = Path(__file__).resolve().parent / "salidas"
        self.salida_dir.var.set(str(default_output))
        self.mapeo_excel.var.set(str(default_output / "mapeo_cuentas_app.xlsx"))

    def append_log(self, text: str):
        self.log.insert(tk.END, text + "\n")
        self.log.see(tk.END)

    def generate(self):
        if not self.maestros.paths:
            messagebox.showwarning("Faltan archivos", "Selecciona al menos un Maestro de Remuneraciones.")
            return
        if not self.gastos.paths:
            messagebox.showwarning("Faltan archivos", "Selecciona al menos un InformeGastoFinanciamiento.")
            return

        self.generate_btn.configure(state="disabled")
        self.status_var.set("Generando centralizacion...")
        self.append_log("Iniciando proceso...")
        threading.Thread(target=self._generate_worker, daemon=True).start()

    def _generate_worker(self):
        try:
            output_dir = Path(self.salida_dir.var.get())
            output_dir.mkdir(parents=True, exist_ok=True)
            area = self.area_var.get().strip().lower()
            period_label = f"{self.month_var.get().strip().upper()} {self.year_var.get().strip()}"
            title_text = f"CENTRALIZACION REMUNERACIONES {self.area_var.get().upper()} {period_label}"
            output_file = output_dir / f"centralizacion_{area}_{period_label.replace(' ', '_').lower()}.xlsx"
            mapping_excel = Path(self.mapeo_excel.var.get().strip() or output_dir / "mapeo_cuentas_app.xlsx")
            mapping_csv = mapping_excel.with_suffix(".csv")

            salida, mapping_csv_path, mapping_xlsx_path = run_pipeline(
                process_files=self.maestros.paths,
                gasto_files=self.gastos.paths,
                central_files=self.centrales.paths,
                cuentas_pdf="",
                salida=str(output_file),
                mapeo=str(mapping_csv),
                mapeo_excel=str(mapping_excel),
                title_text=title_text,
                asiento_files=self.asientos.paths,
            )

            wb = load_workbook(mapping_xlsx_path, data_only=True)
            ws = wb["CompletarAqui"]
            pending = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if not row[6])

            self.after(
                0,
                lambda: self._generation_done(
                    salida=str(salida),
                    mapping_xlsx=str(mapping_xlsx_path),
                    mapping_csv=str(mapping_csv_path),
                    pending=pending,
                ),
            )
        except Exception as exc:
            self.after(0, lambda exc=exc: self._generation_failed(exc))

    def _generation_done(self, salida: str, mapping_xlsx: str, mapping_csv: str, pending: int):
        self.generate_btn.configure(state="normal")
        self.status_var.set("Generacion completada.")
        self.append_log(f"Archivo generado: {salida}")
        self.append_log(f"Mapeo Excel: {mapping_xlsx}")
        self.append_log(f"Mapeo CSV: {mapping_csv}")
        self.append_log(f"Cuentas pendientes: {pending}")
        if pending:
            messagebox.showinfo(
                "Generacion completada",
                f"Se generó la centralizacion.\nQuedaron {pending} cuentas por completar en:\n{mapping_xlsx}",
            )
        else:
            messagebox.showinfo("Generacion completada", f"Centralizacion lista en:\n{salida}")

    def _generation_failed(self, exc: Exception):
        self.generate_btn.configure(state="normal")
        self.status_var.set("Error al generar.")
        self.append_log(f"Error: {exc}")
        messagebox.showerror("Error", str(exc))


if __name__ == "__main__":
    App().mainloop()
